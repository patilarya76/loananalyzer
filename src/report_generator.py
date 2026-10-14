"""
report_generator.py
--------------------
Responsible for Feature 7: Automated Excel MIS Report.

Generates a multi-sheet, executive-ready Excel MIS report named
'collection_mis_report.xlsx' using openpyxl.

Worksheets included:
1. Executive Summary  - Portfolio KPI dashboard cards & overview table.
2. Loan Details        - Granular loan records with payment status & DPD bucket.
3. DPD Analysis        - Aging bucket breakdown with an embedded openpyxl BarChart.
4. Branch Performance  - Branch-wise recovery amounts and collection efficiency.
5. Agency Performance  - Collection agency scorecard.
6. Data Quality Issues - Complete audit log of flagged records and anomaly treatments.

Professional Polish:
- Header fills: Navy Blue (#1F4E79) with bold white text.
- Formats: Currency (₹#,##0), Percentages (0.0%), Integers (#,##0).
- Auto-adjusted column widths.
- Freeze panes and auto-filters.
"""

import os
from typing import Dict, Any
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, Reference


def create_excel_mis_report(
    output_path: str,
    executive_kpis: Dict[str, Any],
    enriched_df: pd.DataFrame,
    dpd_df: pd.DataFrame,
    branch_df: pd.DataFrame,
    agency_df: pd.DataFrame,
    issues_df: pd.DataFrame
) -> str:
    """
    Builds the complete multi-sheet Excel MIS Report.
    """
    os.makedirs(os.path.dirname(output_path), exist_ok=True)

    wb = openpyxl.Workbook()
    # Remove default empty sheet
    wb.remove(wb.active)

    # -------------------------------------------------------------
    # Define Corporate Style Palettes
    # -------------------------------------------------------------
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    kpi_title_font = Font(name="Calibri", size=14, bold=True, color="1F4E79")
    kpi_sub_font = Font(name="Calibri", size=10, italic=True, color="595959")
    card_label_font = Font(name="Calibri", size=10, color="595959")
    card_value_font = Font(name="Calibri", size=16, bold=True, color="1F4E79")
    card_fill = PatternFill(start_color="F2F5F9", end_color="F2F5F9", fill_type="solid")
    zebra_fill = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")

    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9")
    )

    currency_format = '"₹ "#,##0'
    percent_format = '0.0%'
    integer_format = '#,##0'

    # Helper function to style standard table
    def format_worksheet_table(ws, title, start_row, df, is_currency_cols=None, is_pct_cols=None, is_int_cols=None):
        if is_currency_cols is None:
            is_currency_cols = []
        if is_pct_cols is None:
            is_pct_cols = []
        if is_int_cols is None:
            is_int_cols = []

        # Title Block
        ws.cell(row=start_row, column=1, value=title).font = kpi_title_font
        table_header_row = start_row + 2

        # Write Headers
        headers = [col.replace("_", " ").title() for col in df.columns]
        for col_idx, header in enumerate(headers, start=1):
            cell = ws.cell(row=table_header_row, column=col_idx, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border
        ws.row_dimensions[table_header_row].height = 24

        # Write Data
        current_r = table_header_row + 1
        for r_idx, row in df.iterrows():
            row_fill = zebra_fill if (r_idx % 2 == 1) else None
            for col_idx, col_name in enumerate(df.columns, start=1):
                val = row[col_name]
                # Format NaN safely
                if pd.isna(val):
                    val = ""

                cell = ws.cell(row=current_r, column=col_idx, value=val)
                cell.font = Font(name="Calibri", size=10)
                cell.border = thin_border
                if row_fill:
                    cell.fill = row_fill

                # Apply number formatting
                if col_name in is_currency_cols and isinstance(val, (int, float)):
                    cell.number_format = currency_format
                    cell.alignment = Alignment(horizontal="right")
                elif col_name in is_pct_cols and isinstance(val, (int, float)):
                    # If DataFrame pct is e.g. 85.5, convert to 0.855 for Excel's % format
                    cell.value = val / 100.0
                    cell.number_format = percent_format
                    cell.alignment = Alignment(horizontal="right")
                elif col_name in is_int_cols and isinstance(val, (int, float)):
                    cell.number_format = integer_format
                    cell.alignment = Alignment(horizontal="right")
                elif isinstance(val, (int, float)):
                    cell.alignment = Alignment(horizontal="right")
                else:
                    cell.alignment = Alignment(horizontal="left")

            current_r += 1

        # Enable Auto-Filter
        last_col_letter = get_column_letter(len(df.columns))
        ws.auto_filter.ref = f"A{table_header_row}:{last_col_letter}{current_r - 1}"
        # Freeze Panes
        ws.freeze_panes = f"A{table_header_row + 1}"

        # Adjust Column Widths
        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 13)

        return table_header_row, current_r - 1

    # =============================================================
    # Sheet 1: Executive Summary
    # =============================================================
    ws_exec = wb.create_sheet(title="Executive Summary")
    ws_exec.views.sheetView[0].showGridLines = True

    # Title Banner
    ws_exec.cell(row=2, column=2, value="LOANPULSE ANALYTICS — EXECUTIVE COLLECTION MIS REPORT").font = Font(name="Calibri", size=16, bold=True, color="1F4E79")
    ws_exec.cell(row=3, column=2, value="Retail Lending Portfolio Risk & Delinquency Performance Review").font = kpi_sub_font

    # KPI Summary Cards (Rows 5 to 7)
    kpis_to_show = [
        ("Total Active Loans", f"{executive_kpis['Total Active Loans']:,}", "Portfolio Count"),
        ("Total Billing Due", f"₹ {executive_kpis['Total Billing Due Amount']:,.0f}", "Monthly Installments Due"),
        ("Total Collected", f"₹ {executive_kpis['Total Amount Collected']:,.0f}", "Cash Recovered"),
        ("Total Outstanding", f"₹ {executive_kpis['Total Outstanding Amount']:,.0f}", "Overdue / Pending Balance"),
        ("Collection Efficiency", f"{executive_kpis['Overall Collection Efficiency (%)']:.1f}%", "Aggregated Recovery Rate")
    ]

    col_start = 2
    for label, val, desc in kpis_to_show:
        ws_exec.merge_cells(start_row=5, start_column=col_start, end_row=5, end_column=col_start + 1)
        ws_exec.merge_cells(start_row=6, start_column=col_start, end_row=6, end_column=col_start + 1)
        ws_exec.merge_cells(start_row=7, start_column=col_start, end_row=7, end_column=col_start + 1)

        c1 = ws_exec.cell(row=5, column=col_start, value=label)
        c1.font = card_label_font
        c1.alignment = Alignment(horizontal="center", vertical="center")

        c2 = ws_exec.cell(row=6, column=col_start, value=val)
        c2.font = card_value_font
        c2.alignment = Alignment(horizontal="center", vertical="center")

        c3 = ws_exec.cell(row=7, column=col_start, value=desc)
        c3.font = Font(name="Calibri", size=8, italic=True, color="7F7F7F")
        c3.alignment = Alignment(horizontal="center", vertical="center")

        for r in range(5, 8):
            for c in range(col_start, col_start + 2):
                ws_exec.cell(row=r, column=c).fill = card_fill
                ws_exec.cell(row=r, column=c).border = thin_border

        col_start += 3

    # Add Executive KPI Data Table in rows 10-18
    ws_exec.cell(row=10, column=2, value="Key Performance Indicators Summary").font = kpi_title_font
    row_kpi = 12
    ws_exec.cell(row=row_kpi, column=2, value="Portfolio Metric").font = header_font
    ws_exec.cell(row=row_kpi, column=2).fill = header_fill
    ws_exec.cell(row=row_kpi, column=3, value="Reported Value").font = header_font
    ws_exec.cell(row=row_kpi, column=3).fill = header_fill
    ws_exec.row_dimensions[row_kpi].height = 22

    for metric_name, metric_val in executive_kpis.items():
        row_kpi += 1
        ws_exec.cell(row=row_kpi, column=2, value=metric_name).font = Font(name="Calibri", size=10)
        c = ws_exec.cell(row=row_kpi, column=3, value=metric_val)
        c.font = Font(name="Calibri", size=10, bold=True)
        if "Amount" in metric_name or "Portfolio" in metric_name:
            c.number_format = currency_format
        elif "(%)" in metric_name:
            c.value = metric_val / 100.0
            c.number_format = percent_format
        else:
            c.number_format = integer_format
        ws_exec.cell(row=row_kpi, column=2).border = thin_border
        ws_exec.cell(row=row_kpi, column=3).border = thin_border

    # Set column widths for Executive Summary
    ws_exec.column_dimensions["A"].width = 4
    ws_exec.column_dimensions["B"].width = 38
    ws_exec.column_dimensions["C"].width = 24
    for c_idx in range(4, 18):
        ws_exec.column_dimensions[get_column_letter(c_idx)].width = 16

    # =============================================================
    # Sheet 2: Loan Details
    # =============================================================
    ws_loans = wb.create_sheet(title="Loan Details")
    ws_loans.views.sheetView[0].showGridLines = True
    display_loan_df = enriched_df[[
        "loan_id", "customer_name", "branch", "collection_agency",
        "loan_amount", "due_amount", "paid_amount", "outstanding_amount",
        "payment_status", "days_past_due", "dpd_bucket", "overpayment_flag"
    ]].copy()

    format_worksheet_table(
        ws_loans,
        title="Granular Loan Repayment & Delinquency Register",
        start_row=2,
        df=display_loan_df,
        is_currency_cols=["loan_amount", "due_amount", "paid_amount", "outstanding_amount"],
        is_int_cols=["days_past_due"]
    )

    # =============================================================
    # Sheet 3: DPD Analysis
    # =============================================================
    ws_dpd = wb.create_sheet(title="DPD Analysis")
    ws_dpd.views.sheetView[0].showGridLines = True
    header_r, last_r = format_worksheet_table(
        ws_dpd,
        title="Delinquency Aging & DPD Bucket Portfolio Breakdown",
        start_row=2,
        df=dpd_df,
        is_currency_cols=["total_due", "total_collected", "total_outstanding"],
        is_pct_cols=["collection_efficiency_pct", "share_of_outstanding_pct"],
        is_int_cols=["loan_count"]
    )

    # Embed Bar Chart for Outstanding Amount by DPD Bucket
    chart = BarChart()
    chart.type = "col"
    chart.style = 10
    chart.title = "Outstanding Amount by DPD Bucket (₹)"
    chart.y_axis.title = "Outstanding Amount (₹)"
    chart.x_axis.title = "DPD Delinquency Bucket"
    chart.width = 16
    chart.height = 10

    # Total Outstanding is column 5 in DPD table
    data_ref = Reference(ws_dpd, min_col=5, min_row=header_r, max_row=last_r)
    cats_ref = Reference(ws_dpd, min_col=1, min_row=header_r + 1, max_row=last_r)
    chart.add_data(data_ref, titles_from_data=True)
    chart.set_categories(cats_ref)
    chart.legend = None
    ws_dpd.add_chart(chart, f"B{last_r + 3}")

    # =============================================================
    # Sheet 4: Branch Performance
    # =============================================================
    ws_branch = wb.create_sheet(title="Branch Performance")
    ws_branch.views.sheetView[0].showGridLines = True
    format_worksheet_table(
        ws_branch,
        title="Branch-Wise Collection Performance Scorecard",
        start_row=2,
        df=branch_df,
        is_currency_cols=["total_due", "total_collected", "total_outstanding"],
        is_pct_cols=["collection_efficiency_pct"],
        is_int_cols=["loan_count"]
    )

    # =============================================================
    # Sheet 5: Agency Performance
    # =============================================================
    ws_agency = wb.create_sheet(title="Agency Performance")
    ws_agency.views.sheetView[0].showGridLines = True
    format_worksheet_table(
        ws_agency,
        title="External & In-House Collection Agency Scorecard",
        start_row=2,
        df=agency_df,
        is_currency_cols=["total_due", "total_collected", "total_outstanding"],
        is_pct_cols=["collection_efficiency_pct"],
        is_int_cols=["loan_count"]
    )

    # =============================================================
    # Sheet 6: Data Quality Issues
    # =============================================================
    ws_issues = wb.create_sheet(title="Data Quality Issues")
    ws_issues.views.sheetView[0].showGridLines = True
    format_worksheet_table(
        ws_issues,
        title="Data Quality Audit Log — Flagged Records & Resolution Treatments",
        start_row=2,
        df=issues_df,
        is_int_cols=["record_index"]
    )

    wb.save(output_path)
    return output_path


if __name__ == "__main__":
    from data_loader import load_loan_data
    from data_validation import validate_loan_data
    from calculations import enrich_loan_dataset
    from analysis import (
        generate_executive_summary_metrics,
        analyze_dpd_buckets,
        analyze_branch_performance,
        analyze_agency_performance,
    )
    import os

    sample_file = os.path.join(os.path.dirname(__file__), "..", "data", "loan_data.csv")
    out_file = os.path.join(os.path.dirname(__file__), "..", "output", "collection_mis_report.xlsx")

    raw = load_loan_data(sample_file)
    valid, issues, _ = validate_loan_data(raw)
    enriched = enrich_loan_dataset(valid)

    kpis = generate_executive_summary_metrics(enriched)
    dpd_summary = analyze_dpd_buckets(enriched)
    branch_summary = analyze_branch_performance(enriched)
    agency_summary = analyze_agency_performance(enriched)

    res = create_excel_mis_report(
        out_file, kpis, enriched, dpd_summary, branch_summary, agency_summary, issues
    )
    print(f"\n[Test] MIS Report generated at: {res}")
